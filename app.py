from __future__ import annotations

import re
import glob
import threading
import webbrowser
from collections import Counter
from datetime import datetime
from io import BytesIO

from flask import Flask, jsonify, render_template, request, send_file
import os
from kiwipiepy import Kiwi
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait

app = Flask(__name__)
STOPWORDS = set("것 수 등 때 점 곳 및 내 저 제 나 너 우리 당신 이번 오늘 어제 내일 이 그 저 이런 그런 어떤 모든 각각 여러 경우 현재 이후 이전 때문 정도 누구 무엇 어디 왜 어떻게 내용 정보 소개 안내 후기 추천 사용 이용 확인 가능 생각 사진 이미지 제품 상품 가격 구매 방문 블로그 포스팅 글 작성".split())
NON_TOPIC_WORDS = {"블로그", "블로그탭", "포스팅", "게시글", "글", "콘텐츠"}
EXPANSION_RULES = (
    (
        ("법인자본금", "실질자본금", "납입자본금", "자본금"),
        (
            "법인 납입자본금", "실질자본금", "기업진단보고서", "자본금 증빙",
            "법인등기부등본", "건설업 자본금", "건설업 등록기준", "자본금 기준",
            "재무제표 자본금", "건설업 면허등록",
        ),
    ),
    (
        ("연말결산", "결산"),
        (
            "연말결산 재무제표", "결산 전 점검", "결산재무제표", "기업진단보고서",
            "실질자본금", "건설업 실태조사", "건설업 면허 유지", "재무상태표 점검",
            "결산 준비", "건설업 등록기준",
        ),
    ),
    (
        ("실태조사",),
        (
            "건설업 실태조사", "실태조사 준비", "기업진단보고서", "결산재무제표",
            "실질자본금", "기술인력", "공제조합 출자금", "건설업 면허 유지",
            "건설업 등록기준", "결산 전 점검",
        ),
    ),
)
PROTECTED_BUSINESS_TERMS = {
    "건설", "건설업", "건설공사업", "건설면허", "건설업면허", "건설컨설팅",
    "면허", "면허등록", "등록기준", "기업진단", "기업진단보고서",
    "실질자본금", "납입자본금", "기술인력", "공제조합",
}
kiwi = Kiwi()


@app.after_request
def allow_cors(response):
    response.headers['Access-Control-Allow-Origin'] = '*'
    response.headers['Access-Control-Allow-Headers'] = 'Content-Type'
    response.headers['Access-Control-Allow-Methods'] = 'GET, POST, OPTIONS'
    return response


def driver():
    options = Options()
    chrome_binary = os.environ.get("CHROME_BIN")
    if not chrome_binary:
        candidates = glob.glob("/opt/render/.cache/ms-playwright/chromium-*/chrome-linux/chrome")
        chrome_binary = candidates[0] if candidates else None
    if chrome_binary:
        options.binary_location = chrome_binary
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--disable-software-rasterizer")
    options.add_argument("--single-process")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--lang=ko-KR")
    options.add_argument("--disable-blink-features=AutomationControlled")
    return webdriver.Chrome(options=options)


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def merge_equivalent_compounds(counter: Counter) -> Counter:
    """띄어쓰기만 다른 복합명사를 하나의 키워드로 합산한다.

    예: '실질자본금'과 '실질 자본금' → '실질자본금'
    """
    grouped: dict[str, list[tuple[str, int]]] = {}
    for phrase, count in counter.items():
        key = re.sub(r"\s+", "", phrase).lower()
        grouped.setdefault(key, []).append((phrase, count))

    merged = Counter()
    for values in grouped.values():
        # 붙여쓴 원문 표기가 있다면 그것을 우선 노출한다.
        display = next((phrase for phrase, _ in values if " " not in phrase), None)
        if not display:
            display = max(values, key=lambda item: item[1])[0]
        merged[display] = sum(count for _, count in values)
    return merged


def terms(text: str) -> Counter:
    """Kiwi 품사 태그 중 일반명사·고유명사만 사용한다."""
    result = Counter()
    for token in kiwi.tokenize(text):
        if token.tag not in {"NNG", "NNP"}:
            continue
        word = token.form.lower()
        if len(word) >= 2 and word not in STOPWORDS:
            result[word] += 1
    return merge_equivalent_compounds(result)


def phrase_terms(text: str) -> Counter:
    """원문 복합명사와 연속 명사를 2~3어절 문구로 묶어 후보를 만든다."""
    result = Counter()

    # 형태소 분석기가 '세액공제'를 '세액'·'공제'로 나누더라도,
    # 실제 블로그 본문에 쓰인 표기인 '세액공제'를 그대로 유지한다.
    for raw_word in re.findall(r"[가-힣A-Za-z0-9]{2,}", text):
        tokens = kiwi.tokenize(raw_word)
        noun_tokens = [token for token in tokens if token.tag in {"NNG", "NNP"}]
        if not tokens or len(noun_tokens) != len(tokens):
            continue
        if len(noun_tokens) >= 2 or len(raw_word) >= 3:
            result[raw_word] += 1

    run: list[str] = []

    def add_run(words: list[str]) -> None:
        for size in (2, 3):
            for index in range(len(words) - size + 1):
                phrase = " ".join(words[index:index + size])
                if len(phrase.replace(" ", "")) >= 4:
                    result[phrase] += 1

    for token in kiwi.tokenize(text):
        if token.tag in {"NNG", "NNP"} and len(token.form) >= 2:
            run.append(token.form)
        else:
            add_run(run)
            run = []
    add_run(run)
    return merge_equivalent_compounds(result)


def top(counter: Counter, count: int = 12):
    return [{"word": word, "count": value} for word, value in counter.most_common(count)]


def protected_keyword(word: str, main_keyword: str) -> bool:
    normalized = re.sub(r"\s+", "", word).lower()
    main_parts = re.findall(r"[가-힣A-Za-z0-9]{2,}", main_keyword.lower())
    protected = {re.sub(r"\s+", "", item).lower() for item in PROTECTED_BUSINESS_TERMS}
    protected.update(main_parts)
    return any(term and term in normalized for term in protected)


def keyword_core(keyword: str) -> str:
    words = [word for word in re.findall(r"[가-힣A-Za-z0-9]{2,}", keyword) if word not in NON_TOPIC_WORDS]
    return " ".join(words) or keyword


def context_expansions(keyword: str, phrases: Counter) -> list[dict]:
    """입력어와 상위 글의 핵심 복합명사를 이용해 별도 확장 후보를 만든다."""
    evidence = re.sub(r"\s+", "", keyword).lower()
    evidence += " ".join(re.sub(r"\s+", "", word).lower() for word in phrases)
    found: list[dict] = []
    for triggers, suggestions in EXPANSION_RULES:
        matched = [trigger for trigger in triggers if re.sub(r"\s+", "", trigger).lower() in evidence]
        if not matched:
            continue
        for suggestion in suggestions:
            found.append({"word": suggestion, "count": "", "reason": f"{matched[0]} 관련 확장"})
        break

    # 사전 규칙에 없는 키워드도 빈칸이 되지 않도록, 상위 글에서 반복된 복합명사와
    # 입력 메인키워드를 조합한다. 원문에 없던 조합이므로 항상 '확장 추천'으로만 표시한다.
    core = keyword_core(keyword)
    normalized_core = re.sub(r"\s+", "", core).lower()
    existing = {re.sub(r"\s+", "", item["word"]).lower() for item in found}
    for phrase, frequency in phrases.most_common(40):
        normalized = re.sub(r"\s+", "", phrase).lower()
        if (
            len(normalized) < 3
            or normalized_core in normalized
            or normalized in normalized_core
            or phrase in STOPWORDS
        ):
            continue
        candidate = f"{core} {phrase}"
        candidate_normalized = re.sub(r"\s+", "", candidate).lower()
        if candidate_normalized in existing:
            continue
        found.append({"word": candidate, "count": "", "reason": f"상위 글 반복 복합명사 · {phrase}"})
        existing.add(candidate_normalized)
        if len(found) >= 12:
            break
    return found


def recommendation(counter: Counter, keyword: str, phrases: Counter | None = None, phrase_stats: dict | None = None) -> dict:
    """빈도·공통 노출·제목 사용·메인키워드 연결성을 함께 반영해 추천한다."""
    core = keyword_core(keyword)
    normalized_core = re.sub(r"\s+", "", core).lower()
    stats = phrase_stats or {}
    candidates: dict[str, int] = {}

    def add_candidate(word: str, score: int) -> None:
        normalized = re.sub(r"\s+", "", word).lower()
        if not word or normalized == normalized_core or len(normalized) < 4:
            return
        candidates[word] = max(candidates.get(word, 0), score)

    for word, frequency in (phrases or Counter()).items():
        metadata = stats.get(word, {})
        documents = metadata.get("documents", 1)
        title_hits = metadata.get("title_hits", 0)
        normalized = re.sub(r"\s+", "", word).lower()
        score = frequency + documents * 8 + title_hits * 6
        if normalized_core and normalized_core in normalized:
            score += 10
        add_candidate(word, score)

        # 본문에 '소득공제'처럼 핵심 복합명사만 있어도,
        # 입력한 메인키워드와 자연스럽게 조합한 문구를 추가한다.
        if (
            normalized_core
            and normalized_core not in normalized
            and " " not in word
            and documents >= 1
            and len(word) >= 3
        ):
            add_candidate(f"{core} {word}", score + 5)

    ranked_candidates = sorted(candidates.items(), key=lambda item: (-item[1], item[0]))
    related = [{"word": word, "count": score} for word, score in ranked_candidates[:12]]

    selected = {item["word"] for item in related}

    excluded = [
        item for item in top(counter, 40)
        if item["word"] not in selected
        and item["word"] != keyword
        and not protected_keyword(item["word"], keyword)
    ][:15]
    expanded = [
        item for item in context_expansions(keyword, phrases or Counter())
        if item["word"] not in selected
    ][:12]
    return {"main": keyword, "related": related, "expanded": expanded, "excluded": excluded}


def post_text(browser, url: str) -> tuple[str, str]:
    browser.get(url)
    try:
        WebDriverWait(browser, 8).until(lambda d: d.find_elements(By.ID, "mainFrame"))
        browser.switch_to.frame("mainFrame")
    except Exception:
        pass
    title = ""
    for selector in [".se-title-text", ".pcol1 .se-module-text.se-title-text", ".htitle"]:
        elements = browser.find_elements(By.CSS_SELECTOR, selector)
        if elements:
            title = clean(elements[0].text)
            break
    body = ""
    for selector in [".se-main-container", "#postViewArea", "#post-view"]:
        elements = browser.find_elements(By.CSS_SELECTOR, selector)
        if elements:
            body = clean(" ".join(e.text for e in elements))
            break
    browser.switch_to.default_content()
    return title, body


def analyze(keyword: str, urls: list[str]) -> dict:
    """본문 수집과 형태소 분석을 분리해 제한된 서버 메모리를 안정적으로 사용한다."""
    items = [{"url": url, "title": "", "rank": rank} for rank, url in enumerate(urls[:5], 1)]
    browser = None
    try:
        browser = driver()
        for item in items:
            try:
                title, body = post_text(browser, item["url"])
                item["title"] = title
                item["_source_text"] = f"{title} {title} {title} {body}"
            except Exception as error:
                item.update({"main": "분석 실패", "sub": [clean(str(error))]})
    finally:
        if browser is not None:
            browser.quit()

    total = Counter()
    phrases = Counter()
    phrase_documents = Counter()
    phrase_title_hits = Counter()
    for item in items:
        if "main" in item:
            item.pop("_source_text", None)
            continue
        try:
            source_text = item.pop("_source_text", "")
            counts = terms(source_text)
            phrase_counts = phrase_terms(source_text)
            phrases.update(phrase_counts)
            phrase_documents.update(phrase_counts.keys())
            phrase_title_hits.update(phrase_terms(item["title"]).keys())
            display_counts = phrase_counts if phrase_counts else counts
            item["main"] = top(display_counts, 1)[0]["word"] if display_counts else "-"
            item["sub"] = [x["word"] for x in top(display_counts, 11)[1:]]
            total.update(counts)
        except Exception as error:
            item.update({"main": "분석 실패", "sub": [clean(str(error))]})

    return {
        "keyword": keyword,
        "posts": items,
        "keywords": top(phrases if phrases else total, 20),
        "phrases": top(phrases, 30),
        "phrase_stats": {
            phrase: {
                "documents": phrase_documents[phrase],
                "title_hits": phrase_title_hits[phrase],
            }
            for phrase in phrases
        },
        "recommendation": recommendation(
            total,
            keyword,
            phrases,
            {
                phrase: {"documents": phrase_documents[phrase], "title_hits": phrase_title_hits[phrase]}
                for phrase in phrases
            },
        ),
    }

def excel_sheet_title(text: str, used: set[str]) -> str:
    base = re.sub(r'[\\/*?:\[\]]', " ", text).strip() or "분석 결과"
    base = base[:28]
    title = base
    index = 2
    while title in used:
        suffix = f" {index}"
        title = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(title)
    return title


def style_header(cells) -> None:
    fill = PatternFill("solid", fgColor="0B3F6A")
    for cell in cells:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def append_keyword_table(sheet, start_row: int, title: str, items: list[dict]) -> int:
    sheet.cell(start_row, 1, title).font = Font(bold=True, size=13)
    sheet.cell(start_row + 1, 1, "키워드")
    sheet.cell(start_row + 1, 2, "횟수")
    style_header(sheet[start_row + 1][0:2])
    row = start_row + 2
    for item in items:
        sheet.cell(row, 1, item.get("word", ""))
        sheet.cell(row, 2, item.get("count", ""))
        row += 1
    return row + 1


def build_excel(data: dict) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "통합 결과"
    summary["A1"] = "네이버 블로그 키워드 분석 결과"
    summary["A1"].font = Font(bold=True, size=16, color="0B3F6A")
    summary["A2"] = f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    summary["A4"] = "분석 검색어 수"
    summary["B4"] = len(data.get("groups", []))
    style_header(summary["A4":"B4"][0])
    append_keyword_table(summary, 6, "전체 통합 키워드", data.get("total", []))
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 12

    recommendation = data.get("recommendation", {})
    recommendation_sheet = workbook.create_sheet("작성 추천")
    recommendation_sheet["A1"] = "최종 작성 추천"
    recommendation_sheet["A1"].font = Font(bold=True, size=16, color="0B3F6A")
    recommendation_sheet["A3"] = "메인키워드"
    recommendation_sheet["B3"] = recommendation.get("main", "")
    style_header(recommendation_sheet["A3":"B3"][0])
    row = 5
    recommendation_sheet.cell(row, 1, "연관키워드")
    recommendation_sheet.cell(row, 2, "추천 점수")
    style_header(recommendation_sheet[row][0:2])
    row += 1
    for item in recommendation.get("related", []):
        recommendation_sheet.cell(row, 1, item.get("word", ""))
        recommendation_sheet.cell(row, 2, item.get("count", ""))
        row += 1
    row += 1
    row += 1
    recommendation_sheet.cell(row, 1, "연관도 기반 확장 추천")
    recommendation_sheet.cell(row, 2, "추천 근거")
    style_header(recommendation_sheet[row][0:2])
    row += 1
    for item in recommendation.get("expanded", []):
        recommendation_sheet.cell(row, 1, item.get("word", ""))
        recommendation_sheet.cell(row, 2, item.get("reason", ""))
        row += 1
    row += 1
    recommendation_sheet.cell(row, 1, "제외 검토 키워드")
    recommendation_sheet.cell(row, 2, "빈도")
    style_header(recommendation_sheet[row][0:2])
    row += 1
    for item in recommendation.get("excluded", []):
        recommendation_sheet.cell(row, 1, item.get("word", ""))
        recommendation_sheet.cell(row, 2, item.get("count", ""))
        row += 1
    recommendation_sheet.column_dimensions["A"].width = 30
    recommendation_sheet.column_dimensions["B"].width = 12

    used_names = {"통합 결과", "작성 추천"}
    for group in data.get("groups", []):
        sheet = workbook.create_sheet(excel_sheet_title(group.get("keyword", "분석 결과"), used_names))
        sheet["A1"] = f"{group.get('keyword', '')} 분석 결과"
        sheet["A1"].font = Font(bold=True, size=16, color="0B3F6A")
        row = append_keyword_table(sheet, 3, "키워드 빈도", group.get("keywords", []))
        sheet.cell(row, 1, "순위")
        sheet.cell(row, 2, "콘텐츠 제목")
        sheet.cell(row, 3, "URL")
        sheet.cell(row, 4, "메인키워드")
        sheet.cell(row, 5, "서브키워드")
        style_header(sheet[row][0:5])
        row += 1
        for post in group.get("posts", []):
            sheet.cell(row, 1, post.get("rank", ""))
            sheet.cell(row, 2, post.get("title", ""))
            sheet.cell(row, 3, post.get("url", ""))
            sheet.cell(row, 4, post.get("main", ""))
            sheet.cell(row, 5, ", ".join(post.get("sub", [])))
            row += 1
        for column, width in {"A": 10, "B": 42, "C": 55, "D": 20, "E": 55}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A4"

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@app.post("/api/analyze")
def api_analyze():
    groups_input = request.json.get("groups", [])[:10]
    groups_input = [g for g in groups_input if g.get("keyword") and g.get("urls")]
    if not groups_input:
        return jsonify(error="키워드와 블로그 URL을 입력해 주세요."), 400
    try:
        groups = [analyze(g["keyword"].strip(), g["urls"]) for g in groups_input]
        combined = Counter()
        combined_phrases = Counter()
        combined_phrase_stats: dict[str, dict[str, int]] = {}
        for group in groups:
            for item in group["keywords"]:
                combined[item["word"]] += item["count"]
            for item in group["phrases"]:
                combined_phrases[item["word"]] += item["count"]
            for phrase, metadata in group["phrase_stats"].items():
                current = combined_phrase_stats.setdefault(phrase, {"documents": 0, "title_hits": 0})
                current["documents"] += metadata["documents"]
                current["title_hits"] += metadata["title_hits"]
        final_recommendation = recommendation(
            combined,
            groups[0]["keyword"],
            combined_phrases,
            combined_phrase_stats,
        )
        return jsonify(
            groups=groups,
            total=top(combined, 25),
            recommendation=final_recommendation,
        )
    except Exception as error:
        message = clean(str(error))
        if not message or message == "Message:":
            message = "크롬에서 네이버 검색 결과를 불러오지 못했습니다. Chrome을 최신 버전으로 업데이트한 뒤 다시 시도해 주세요."
        return jsonify(error=message), 500


@app.post("/api/export")
def api_export():
    data = request.json or {}
    if not data.get("groups"):
        return jsonify(error="먼저 분석을 완료해 주세요."), 400
    output = build_excel(data)
    filename = f"naver_blog_keyword_analysis_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/")
@app.get("/health")
def health():
    return jsonify(status="ok")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
